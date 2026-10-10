"""
app/services/docx_loader.py
────────────────────────────
Extracts plain text from a .docx file-like object entirely in memory.

DOCX has no fixed pagination (page breaks depend on the renderer, font
substitution, etc.), so unlike pdf_loader there is no per-page offset map —
the document is treated as a single flowing document for citation
highlighting (the frontend highlights directly against rendered HTML rather
than a page number).
"""

from dataclasses import dataclass, field

import docx  # python-docx

from app.services.pdf_loader import ExtractResult, PageOffset


def extract_docx(file) -> ExtractResult:
    """
    Extract text from a .docx file-like object (in-memory).

    Paragraphs are joined with "\\n\\n" (mirroring pdf_loader's page join) so
    downstream chunking/quote-verification code can treat both formats the
    same way. Table cell text is included (common in contracts — signature
    blocks, fee schedules) since skipping tables would silently drop real
    document content.

    Returns a single PageOffset covering the whole document (page_number=1)
    so document_pages stays populated for both formats without the rest of
    the pipeline needing format-specific branching.
    """
    raw_bytes = file.read() if not isinstance(file, (bytes, bytearray)) else file

    import io
    document = docx.Document(io.BytesIO(raw_bytes))

    parts: list[str] = []
    for para in document.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    text = "\n\n".join(parts)
    pages = [PageOffset(page_number=1, char_start=0, char_end=len(text))] if text else []
    return ExtractResult(text=text, pages=pages)
